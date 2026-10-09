"""CBF(통합 입찰·낙찰 DB) 로더와 데이터 점검.

CBF.xlsx 의 '통합데이터' 시트(V94/V95 형식)를 분석용 표준 컬럼으로 바꾼다.

표준 컬럼
--------
notice, org, org_key, industry, industry_group, region, est_price, base, lower_rate, a_value,
net_cost, band, open_dt, date, expected, rate(실제 사정율 0기준 %), floor(낙찰하한가),
winner_bid, winner_rate(1순위 가정 사정율), n_bidders, status(COMPLETED/PENDING/OTHER),
award_method, lower_limit_status, quality

실데이터 확인 사항(2025-01 ~ 2026-09, 전기·통신·소방 3,552건)
--------------------------------------------------------------
* 사정율은 '예가/기초(0%)'(소수 4자리 절사) 대신 예정가격/기초금액 으로 정확히 다시 계산한다.
* LH 일부 공고는 원본 사정율이 (예정가격-기초금액)/(기초금액-A값) 기준이다 -> 자동 감지(rate_scale).
* 1순위사정율(0%) = 1순위 투찰금액이 유효한 사정율 상한 = min(하한가 역산, 순공사원가 98% 역산).
* 순공사원가 98% 기준은 '사정율을 반영한' 금액으로 적용된다(verify_net_cost_rule).
* 낙찰하한가 원단위 절상은 CBF 파이프라인 계산과 일치한다(내부 일관성). 1순위가 '절상값-1원'에 놓인 사례도 없지만,
  1순위가 하한가 1원 이내에 놓이는 일이 드물어 이 독립점검은 검정력이 낮다(verify_floor_formula).
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable, Mapping

import numpy as np
import pandas as pd

from .bid import implied_rate_array

# 원본 컬럼 후보(앞에 있을수록 우선)
_COLS: dict[str, tuple[str, ...]] = {
    "notice": ("공고번호",),
    "org": ("발주기관",),
    "industry": ("업종",),
    "region": ("지역",),
    "est_price": ("추정가격",),
    "base": ("기초금액",),
    "lower_rate": ("낙찰하한율", "낙찰하한율_정규화값"),
    "a_value": ("A값", "A값_계산용"),
    "net_cost": ("순공사원가",),
    "band_text": ("예가변동폭",),
    "band_low": ("예가변동폭_low",),
    "band_high": ("예가변동폭_high",),
    "open_dt": ("개찰일시", "개찰일"),
    "open_date": ("개찰일",),
    "expected": ("예정가격",),
    "rate": ("예가/기초(0%)", "예정가격사정율_계산(0%)"),
    "rate100": ("예가/기초(100%)",),
    "floor": ("낙찰하한가",),
    "winner_bid": ("1순위투찰금액",),
    "winner_rate_src": ("1순위사정율(0%)",),
    "n_bidders": ("업체수",),
    "record_status": ("record_status",),
    "quality": ("quality_status", "분석품질상태"),
    "award_method": ("낙찰자선정방법_표준",),
    "lower_limit_status": ("낙찰하한율상태", "lower_limit_status"),
    "contract_type": ("계약유형",),
}

_NUMERIC = ("est_price", "base", "lower_rate", "a_value", "net_cost", "band_low", "band_high",
            "expected", "rate", "rate100", "floor", "winner_bid", "winner_rate_src", "n_bidders")

#: 발주기관명 표기 오류 보정(정규화 후 적용)
ORG_ALIASES: dict[str, str] = {
    "조달청대전지방지방조달청": "조달청대전지방조달청",
}


def org_key(name: object, aliases: Mapping[str, str] = ORG_ALIASES) -> str:
    """발주기관 비교용 키: 공백 제거 + '지방지방' 같은 반복 오타 보정 + 별칭."""
    s = re.sub(r"\s+", "", str(name or ""))
    s = s.replace("지방지방", "지방")
    return aliases.get(s, s)


def industry_group(text: object) -> str:
    s = str(text or "")
    has_comm, has_elec, has_fire = "통신" in s, "전기" in s, "소방" in s
    if has_fire:
        return "소방"
    if has_comm and has_elec:
        return "전기+통신"
    if has_comm:
        return "통신"
    if has_elec:
        return "전기"
    return "기타"


def parse_band(text: object, low: float = np.nan, high: float = np.nan) -> float:
    """'-3/+3', '±2', '-2.5/+2.5' -> 3.0 / 2.0 / 2.5."""
    if np.isfinite(low) and np.isfinite(high):
        return float(max(abs(low), abs(high)))
    nums = [abs(float(v)) for v in re.findall(r"\d+(?:\.\d+)?", str(text or ""))]
    return float(max(nums)) if nums else np.nan


def _first(df: pd.DataFrame, names: Iterable[str]) -> pd.Series:
    for n in names:
        if n in df.columns:
            return df[n]
    return pd.Series(np.nan, index=df.index)


def read_cbf_excel(path: str | Path, sheet: str = "통합데이터") -> pd.DataFrame:
    """CBF.xlsx 의 한 시트를 원본 그대로 읽는다(read_only, 값만)."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            return pd.DataFrame()
        cols = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(header)]
        rows = [r for r in it if r is not None and any(v is not None for v in r)]
    finally:
        wb.close()
    return pd.DataFrame(rows, columns=cols)


def standardize_cbf(raw: pd.DataFrame) -> pd.DataFrame:
    """원본 CBF 표 -> 표준 컬럼 표."""
    out = pd.DataFrame(index=raw.index)
    for k, names in _COLS.items():
        out[k] = _first(raw, names)
    for k in _NUMERIC:
        col = out[k]
        if col.dtype == object:  # '61,795,000' 같은 천단위 구분 문자열
            col = col.astype(str).str.replace(",", "", regex=False).str.replace("%", "", regex=False).str.strip()
        out[k] = pd.to_numeric(col, errors="coerce")
    out["notice"] = out["notice"].astype(str).str.strip()
    out["org"] = out["org"].astype(str).str.strip()
    out["org_key"] = out["org"].map(org_key)
    out["industry"] = out["industry"].fillna("").astype(str)
    out["industry_group"] = out["industry"].map(industry_group)
    out["open_dt"] = pd.to_datetime(out["open_dt"], errors="coerce")
    od = pd.to_datetime(out["open_date"], errors="coerce")
    out["open_dt"] = out["open_dt"].fillna(od)
    out["date"] = out["open_dt"].dt.normalize()
    src = out["rate"].fillna(out["rate100"] - 100.0)
    exact = (out["expected"] / out["base"] - 1.0) * 100.0
    # 원본 사정율이 (기초금액-A값) 기준인지 감지: 원본 x (B-A)/B = 정확값
    k = (out["base"] - out["a_value"].fillna(0.0)) / out["base"]
    alt = (src.abs() > 0.02) & ((exact - src).abs() > 2e-4) & ((src * k - exact).abs() <= 2e-4) & (k < 0.9995)
    out["rate_scale"] = np.where(alt, "base_minus_a", "base")
    out["rate_src"] = src
    out["rate"] = exact.where(exact.notna(), np.where(alt, src * k, src))
    out["winner_rate_src"] = np.where(alt, out["winner_rate_src"] * k, out["winner_rate_src"])
    out["band"] = [parse_band(t, lo, hi) for t, lo, hi in zip(out["band_text"], out["band_low"], out["band_high"])]
    # 1순위 유효 상한 사정율: 투찰금액에서 직접 역산(순공사원가 기준 포함, 원본 값과 교차검증)
    nc_applies = out["net_cost"].where(out["est_price"].fillna(0) < 1e10)
    out["winner_rate"] = implied_rate_array(out["base"], out["winner_bid"], out["lower_rate"], out["a_value"], nc_applies)
    out["winner_rate_diff"] = (out["winner_rate"] - out["winner_rate_src"]).abs()
    out["gap"] = out["winner_rate"] - out["rate"]
    floor_int = np.ceil(out["floor"] - 1e-6)
    out["winner_ge_floor"] = out["winner_bid"] >= floor_int
    rs = out["record_status"].fillna("").astype(str).str.upper()
    done = out["rate"].notna() & out["winner_bid"].notna()
    out["status"] = np.where(rs.eq("COMPLETED") | (rs.eq("") & done), "COMPLETED",
                             np.where(rs.eq("PENDING") | (rs.eq("") & out["rate"].isna()), "PENDING", "OTHER"))
    out = out.drop(columns=["open_date", "rate100", "band_low", "band_high"])
    return out.sort_values(["open_dt", "notice"], kind="mergesort").reset_index(drop=True)


def load_cbf(path: str | Path, sheet: str = "통합데이터") -> pd.DataFrame:
    """CBF 파일(xlsx/csv/parquet) -> 표준 컬럼 표."""
    p = Path(path)
    if p.suffix.lower() in (".csv", ".txt"):
        try:
            raw = pd.read_csv(p, encoding="utf-8-sig", low_memory=False, thousands=",")
        except UnicodeDecodeError:  # 한글 Windows 기본(cp949)으로 저장된 CSV
            raw = pd.read_csv(p, encoding="cp949", low_memory=False, thousands=",")
    elif p.suffix.lower() == ".parquet":
        raw = pd.read_parquet(p)
    else:
        raw = read_cbf_excel(p, sheet)
    return standardize_cbf(raw)


# ---------------------------------------------------------------------------
# 분석 대상 선별
# ---------------------------------------------------------------------------
def eval_sample_mask(df: pd.DataFrame, *, max_rate_diff: float = 1e-3) -> pd.Series:
    """백테스트 평가 대상: 개찰 전에 알 수 있는 조건 + 데이터 일관성만으로 고른다(결과로 거르지 않음).

    완료, 고정 하한율, 예가변동폭 확인, 업체수 >= 2, 1순위 금액이 하한가 이상,
    1순위 역산 사정율이 원본과 일치(원본이 없으면 생략).
    """
    lls = df["lower_limit_status"].fillna("FIXED_LOWER_LIMIT").astype(str)
    consistent = df["winner_rate_diff"].le(max_rate_diff) | df["winner_rate_src"].isna()
    floor_ok = df["winner_ge_floor"] | df["floor"].isna()
    # 순공사원가 98%(사정율 반영) 기준을 1순위가 어긴 공고는 자료 불일치로 본다(그 기준에서는 어떤 전략도 질 수밖에 없음)
    nc98 = 0.98 * df["net_cost"] * (1 + df["rate"] / 100.0)
    nc_ok = ~(df["net_cost"].gt(0) & df["est_price"].fillna(0).lt(1e10) & (df["winner_bid"] < np.ceil(nc98 - 1e-6)))
    return (
        df["status"].eq("COMPLETED")
        & lls.eq("FIXED_LOWER_LIMIT")
        & df["band"].notna()
        & df["rate"].notna()
        & df["winner_bid"].notna()
        & df["n_bidders"].ge(2)
        & floor_ok
        & nc_ok
        & consistent
    )


def gap_sample_mask(df: pd.DataFrame, *, max_gap: float | None = None, max_rate_diff: float = 1e-3) -> pd.Series:
    """경쟁사 모형 학습 표본 = 평가 대상 중 1순위가 실제 사정율 위에 있는(간격 > 0) 공고.

    max_gap 을 주면 그보다 큰 간격을 제외한다(기본: 제외하지 않음. 예가변동폭 밖 1순위는 적합 단계에서 제외).
    """
    m = eval_sample_mask(df, max_rate_diff=max_rate_diff) & df["gap"].gt(0)
    if max_gap is not None:
        m &= df["gap"].le(max_gap)
    return m


# ---------------------------------------------------------------------------
# 산식·데이터 검증
# ---------------------------------------------------------------------------
def verify_floor_formula(df: pd.DataFrame) -> pd.DataFrame:
    """낙찰하한가 원단위 규칙 검증: 정수로 저장된 하한가와 CEIL/HALF_UP/FLOOR 계산값 일치율."""
    from .bid import AwardRule, lower_limit_from_expected

    m = df["expected"].notna() & df["lower_rate"].notna() & df["floor"].notna()
    m &= (df["floor"] % 1 == 0)
    rows = []
    for mode in ("CEIL", "HALF_UP", "FLOOR"):
        hits = 0
        for p, a, r, f in zip(df.loc[m, "expected"], df.loc[m, "a_value"], df.loc[m, "lower_rate"], df.loc[m, "floor"]):
            rule = AwardRule(float(r), float(a) if pd.notna(a) else 0.0, rounding=mode)
            hits += int(lower_limit_from_expected(float(p), rule) == int(f))
        n = int(m.sum())
        rows.append({"rounding": mode, "n": n, "exact_match": hits / n if n else np.nan})
    rows.append({"rounding": "(원본 하한가가 소수로 저장되어 제외)", "n": int((df["floor"].notna() & (df["floor"] % 1 != 0)).sum()), "exact_match": np.nan})
    out = pd.DataFrame(rows)
    out["count"] = np.nan
    out["note"] = ["CBF 낙찰하한가 컬럼과의 일치(내부 일관성 점검)", "", "", ""]
    # 독립 근거: 하한가 계산값이 소수인 공고에서 1순위가 '절상값' 또는 '절상값-1원'에 놓였는지
    calc = []
    for p, a, r in zip(df["expected"], df["a_value"], df["lower_rate"]):
        if not (np.isfinite(p) and np.isfinite(r)):
            calc.append(np.nan)
            continue
        aa = float(a) if np.isfinite(a) else 0.0
        calc.append(((p - aa) * r / 100.0 + aa) if aa > 0 else p * r / 100.0)
    calc = pd.Series(calc, index=df.index)
    frac = calc.notna() & ((calc % 1) > 1e-6) & ((calc % 1) < 1 - 1e-6) & df["winner_bid"].notna()
    ceil_v = np.ceil(calc)
    at_ceil = int((frac & (df["winner_bid"] == ceil_v)).sum())
    at_ceil_m1 = int((frac & (df["winner_bid"] == ceil_v - 1)).sum())
    indep = pd.DataFrame([{
        "rounding": "독립점검: 1순위 금액 = 절상 하한가", "n": int(frac.sum()), "exact_match": np.nan, "count": at_ceil,
        "note": "1순위가 정확히 절상 하한가에 투찰한 건수(절상 규칙이면 유효)"},
        {"rounding": "독립점검: 1순위 금액 = 절상 하한가 - 1원", "n": int(frac.sum()), "exact_match": np.nan, "count": at_ceil_m1,
         "note": "0 이면 절사/반올림 규칙의 흔적 없음. 단 1순위가 하한가 1원 이내인 일이 드물어(기대 1건 미만) 검정력이 낮은 보조 근거"}])
    return pd.concat([out, indep], ignore_index=True)


def verify_net_cost_rule(df: pd.DataFrame) -> pd.DataFrame:
    """순공사원가 98% 기준이 사정율 반영인지 검증.

    순공사원가 98% 가 하한가보다 높은 공고에서 1순위가 그 기준을 지켰는지 본다.
    실제 적용 기준이라면 1순위(유효 최저가)는 거의 항상 기준 이상이어야 한다.
    """
    m = (df["status"].eq("COMPLETED") & df["net_cost"].gt(0) & df["winner_bid"].notna()
         & df["floor"].notna() & df["est_price"].lt(1e10))
    d = df[m]
    rows = []
    for label, nc98 in (("공고금액 그대로(미반영)", 0.98 * d["net_cost"]),
                        ("사정율 반영", 0.98 * d["net_cost"] * (1 + d["rate"] / 100.0))):
        binding = nc98 > d["floor"]
        n = int(binding.sum())
        ok = int((d.loc[binding, "winner_bid"] >= nc98[binding]).sum())
        rows.append({"rule": label, "cases_where_rule_binds": n, "winner_complies": ok,
                     "comply_share": ok / n if n else np.nan,
                     "all_rows_winner_below_rule": int((d["winner_bid"] < nc98).sum())})
    return pd.DataFrame(rows)


def quality_report(df: pd.DataFrame) -> pd.DataFrame:
    """분석에 영향을 주는 데이터 문제 집계."""
    c = df["status"].eq("COMPLETED")
    items = [
        ("전체 공고", len(df)),
        ("완료(개찰결과 있음)", int(c.sum())),
        ("개찰 전(PENDING)", int(df["status"].eq("PENDING").sum())),
        ("기타(입찰/낙찰 한쪽만 등)", int(df["status"].eq("OTHER").sum())),
        ("완료 중 하한율 없음", int((c & df["lower_rate"].isna()).sum())),
        ("완료 중 예가변동폭 없음", int((c & df["band"].isna()).sum())),
        ("1순위 역산 사정율 ≠ 원본(>0.001%p)", int((c & df["winner_rate_diff"].gt(1e-3)).sum())),
        ("1순위가 실제 사정율보다 낮음(간격<0)", int((c & df["gap"].lt(0)).sum())),
        ("1순위 금액 < 낙찰하한가(자료 이상)", int((c & ~df["winner_ge_floor"] & df["floor"].notna()).sum())),
        ("간격 > 3%p", int((c & df["gap"].gt(3)).sum())),
        ("원본 사정율이 (기초금액-A값) 기준(LH 등)", int((df["rate_scale"] == "base_minus_a").sum())),
        ("순공사원가 기준이 1순위 사정율을 결정", int((c & df["net_cost"].gt(0) & df["winner_rate"].lt(
            implied_rate_array(df["base"], df["winner_bid"], df["lower_rate"], df["a_value"]) - 1e-9)).sum())),
        ("업체수 < 2", int((c & df["n_bidders"].lt(2)).sum())),
        ("발주기관명 '지방지방' 오타", int(df["org"].str.contains("지방지방", na=False).sum())),
        ("백테스트 평가 가능(결과로 거르지 않음)", int(eval_sample_mask(df).sum())),
        ("경쟁사 모델 학습 가능(간격 > 0)", int(gap_sample_mask(df).sum())),
    ]
    return pd.DataFrame(items, columns=["항목", "건수"])
